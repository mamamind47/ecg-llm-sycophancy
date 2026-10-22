"""Agreement between the hand labels (share/ECG_labels.xlsx, completed) and the LLM judge, plus the arm-a rationale check.

Usage: python3 scripts/20_label_agreement.py COMPLETED.xlsx
Writes analysis/label_agreement.md and analysis/label_disagreements.csv (every judge/human mismatch, for adjudication).
"""
import json
import pathlib
import sys
from collections import Counter

import pandas as pd
from openpyxl import load_workbook

FIELDS = ["main_dx", "stemi_claimed", "agrees_with_user", "refers_to_clinician", "confidence_hedged"]


def kappa(a, b):
    n = len(a)
    po = sum(x == y for x, y in zip(a, b)) / n
    ca, cb = Counter(a), Counter(b)
    pe = sum(ca[k] * cb[k] for k in set(a) | set(b)) / n ** 2
    return (po - pe) / (1 - pe) if pe < 1 else 1.0, po


def main(path):
    key = pd.read_csv("analysis/label_key.csv").set_index("id")
    wb = load_workbook(path)
    out, dis = ["# Hand labels vs LLM judge\n"], []
    for sheet in [s for s in wb.sheetnames if s.startswith("arm r")]:
        rows = []
        for r in wb[sheet].iter_rows(min_row=2, values_only=True):
            src = pathlib.Path(key.loc[r[0], "file"])
            j = json.loads((pathlib.Path("data/hints/r_judged") / src.parent.name / src.name).read_text())["verdict"]
            rows.append({"id": r[0], "model": src.parent.name, "file": str(src), "message": r[1], "reply": r[2],
                         **{f"h_{f}": str(v).strip() for f, v in zip(FIELDS, r[3:8])},
                         **{f"j_{f}": str(j.get(f)).strip() for f in FIELDS}, "note": r[8]})
        df = pd.DataFrame(rows)
        out.append(f"## {sheet} (n = {len(df)})\n\n| Field | Cohen's κ | Raw agreement |\n|---|---|---|")
        for f in FIELDS:
            k, po = kappa(list(df[f"j_{f}"]), list(df[f"h_{f}"]))
            out.append(f"| {f} | {k:.2f} | {po:.0%} |")
            for _, r in df[df[f"j_{f}"] != df[f"h_{f}"]].iterrows():
                dis.append({"sheet": sheet, "id": r.id, "field": f, "judge": r[f"j_{f}"], "human": r[f"h_{f}"],
                            "model": r.model, "message": r.message, "note": r.note, "reply": r.reply})
        out.append("")
    a = list(wb["arm a"].iter_rows(min_row=2, values_only=True))
    for i, name in ((4, "diagnosis matches rationale"), (5, "STEMI matches rationale")):
        c = Counter(str(r[i]).strip() for r in a)
        out.append(f"- arm a, {name}: " + ", ".join(f"{k} {v}" for k, v in c.most_common()) + f" (n = {len(a)})")
    pd.DataFrame(dis).to_csv("analysis/label_disagreements.csv", index=False)
    pathlib.Path("analysis/label_agreement.md").write_text("\n".join(out) + "\n")
    print("\n".join(out))
    print(f"\n{len(dis)} field-level disagreements -> analysis/label_disagreements.csv")


if __name__ == "__main__":
    main(sys.argv[1])
