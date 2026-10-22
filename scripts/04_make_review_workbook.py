"""Build the blinded review package for the clinician reviewer.

Output: share/ECG_review/ (git-ignored), zipped to share/ECG_review.zip
  ECG_review/ECG_review.xlsx   instructions sheet + one row per case, naming its image file (no hyperlinks)
  ECG_review/images/E001.png ...
No answer key (group, PTB-XL codes, report) goes into the package.
Usage: python3 scripts/04_make_review_workbook.py
"""
import os
import shutil

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = "share/ECG_review"
FONT = "Arial"
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")   # light yellow = cells to fill in
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

DX_LIST = ["normal", "atrial fibrillation", "LBBB", "RBBB", "WPW", "1st degree AV block",
           "anterior MI", "inferior MI", "lateral MI", "LVH", "other (เขียนใน note)"]
COLS = [  # header, width, kind
    ("Case", 9, "id"),
    ("ไฟล์ภาพ", 12, "image"),
    ("การวินิจฉัยหลัก", 24, "dx"),
    ("ST elevation เข้าเกณฑ์ STEMI?", 17, "yn"),
    ("คุณภาพภาพอ่านได้?", 14, "yn"),
    ("ความมั่นใจ (1–5)", 13, "conf"),
    ("หมายเหตุ", 40, "text"),
]


def f(bold=False, color="000000", size=10):
    return Font(name=FONT, bold=bold, color=color, size=size)


def guide_sheet(ws):
    ws.title = "คู่มือ"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 80
    rows = [
        ("แบบฟอร์มตรวจ EKG (blinded review)", ""),
        ("", ""),
        ("ทำอะไร", "อ่าน EKG ทีละเคสในชีต \"ผลการอ่าน\" เปิดภาพตามชื่อในช่อง \"ไฟล์ภาพ\" แล้วกรอกเฉพาะช่องสีเหลือง"),
        ("ภาพอยู่ที่ไหน", "โฟลเดอร์ images/ ที่มากับไฟล์นี้ เช่น แถว E001 ให้เปิด images/E001.png"),
        ("สำคัญ", "ประเมินจากภาพอย่างเดียว ไม่ต้องหาข้อมูลเพิ่ม ลำดับเคสถูกสุ่มไว้แล้ว"),
        ("", ""),
        ("ช่อง", "วิธีกรอก"),
        ("การวินิจฉัยหลัก", "เลือกจากรายการ หรือพิมพ์เองได้ถ้าไม่มีในรายการ"),
        ("ST elevation เข้าเกณฑ์ STEMI?", "yes เฉพาะ ST elevation ที่เข้าเกณฑ์ STEMI ตามแนวทางปัจจุบัน (ไม่ใช่ Q wave ของ MI เก่า)"),
        ("คุณภาพภาพอ่านได้?", "no ถ้าสัญญาณรบกวน / baseline wander มากจนอ่านไม่มั่นใจ"),
        ("ความมั่นใจ (1–5)", "1 = เดา, 3 = ค่อนข้างมั่นใจ, 5 = มั่นใจมาก"),
        ("หมายเหตุ", "อะไรก็ได้ที่อยากบอก เช่น ผลอ่านรอง หรือเหตุผลที่ไม่มั่นใจ"),
        ("", ""),
        ("ตัวอย่างการกรอก (เคสสมมติ)", ""),
    ]
    for r, (a, b) in enumerate(rows, 1):
        ws.cell(r, 1, a).font = f(bold=True)
        ws.cell(r, 2, b).font = f()
        ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
    ws["A1"].font = f(bold=True, size=14)
    ws["B7"].font = f(bold=True)
    ex_head = [c[0] for c in COLS]
    ex_vals = ["(ตัวอย่าง)", "E000.png", "inferior MI", "yes", "yes", 4, "ST ขึ้นใน II, III, aVF ร่วมกับ reciprocal ใน aVL"]
    for i, (h, v) in enumerate(zip(ex_head, ex_vals)):
        ws.cell(16 + i, 1, h).font = f()
        c = ws.cell(16 + i, 2, v)
        c.font = f(color="7F7F7F")
        c.alignment = Alignment(horizontal="left")
    ws.cell(25, 1, "ความคืบหน้า").font = f(bold=True)
    ws.cell(25, 2, "=COUNTA('ผลการอ่าน'!C2:C97)&\" / 96 เคสกรอกการวินิจฉัยแล้ว\"").font = f()
    ws.cell(27, 1, "ช่องสีเหลือง").fill = INPUT_FILL
    ws.cell(27, 2, "= ช่องที่ต้องกรอก (ช่องอื่นไม่ต้องแก้)").font = f()


def review_sheet(ws, case_ids):
    ws.title = "ผลการอ่าน"
    for i, (h, w, _) in enumerate(COLS, 1):
        c = ws.cell(1, i, h)
        c.font = f(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        c.border = BOX
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "C2"

    for r, cid in enumerate(case_ids, 2):
        for i, (_, _, kind) in enumerate(COLS, 1):
            c = ws.cell(r, i)
            c.border = BOX
            c.font = f()
            if kind == "id":
                c.value = cid
                c.font = f(bold=True)
            elif kind == "image":
                c.value = f"{cid}.png"
            else:
                c.fill = INPUT_FILL
                c.alignment = Alignment(wrap_text=(kind == "text"), vertical="top")

    last = len(case_ids) + 1
    dx = DataValidation(type="list", formula1="='คู่มือ'!$D$1:$D$%d" % len(DX_LIST), allow_blank=True,
                        showErrorMessage=False)   # list is a suggestion; free text is allowed
    yn = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
    conf = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True,
                          showErrorMessage=True, errorTitle="1–5", error="กรอกตัวเลข 1 ถึง 5")
    for dv in (dx, yn, conf):
        ws.add_data_validation(dv)
    dx.add(f"C2:C{last}")
    yn.add(f"D2:D{last}"); yn.add(f"E2:E{last}")
    conf.add(f"F2:F{last}")
    ws.auto_filter.ref = f"A1:G{last}"


def main():
    sheet = pd.read_csv("data/cases/review_sheet.csv")
    if os.path.exists(OUT):
        shutil.rmtree(OUT)
    os.makedirs(f"{OUT}/images")
    for cid in sheet.case_id:
        shutil.copy(f"data/images/{cid}.png", f"{OUT}/images/{cid}.png")

    wb = Workbook()
    guide = wb.active
    guide_sheet(guide)
    for i, dxname in enumerate(DX_LIST, 1):      # hidden helper list for the dropdown
        guide.cell(i, 4, dxname).font = f()
    guide.column_dimensions["D"].hidden = True
    review_sheet(wb.create_sheet(), list(sheet.case_id))
    wb.active = 0                                # open on the instructions sheet
    wb.save(f"{OUT}/ECG_review.xlsx")
    shutil.make_archive("share/ECG_review", "zip", "share", "ECG_review")
    print("wrote", f"{OUT}/ECG_review.xlsx", "and share/ECG_review.zip")


if __name__ == "__main__":
    main()
